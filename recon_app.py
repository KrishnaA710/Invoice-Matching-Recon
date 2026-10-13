import streamlit as st
import pandas as pd
import numpy as np
import openpyxl
import xlsxwriter
import zipfile
import re
import io
import os
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl import Workbook

st.set_page_config(page_title="LNG Invoice Reconciliation", page_icon="📊", layout="wide")

st.title("📊 LNG Invoice Reconciliation Tool")
st.caption("Upload Topaz + Allegro files to generate the full recon workbook")

# ── Sidebar settings ──
with st.sidebar:
    st.header("Settings")
    book = st.selectbox("Book", ["LSC","LSA","LSY","LSS","LSP","LSF","LSX","LJC","LLS","Other"])
    if book == "Other":
        book = st.text_input("Custom book name")
    env = st.selectbox("Environment", ["UAT1","DEV3","PROD"])
    cutoff_str = st.text_input("Cutoff date", "1 June 2026")
    cutoff = pd.Timestamp("2026-06-01")

st.markdown("---")

# ── File upload ──
col1, col2 = st.columns(2)
with col1:
    st.subheader("Topaz extract")
    topaz_file = st.file_uploader("Upload Topaz .xlsx", type=["xlsx","xls"], key="topaz")
with col2:
    st.subheader("Allegro extract")
    allegro_file = st.file_uploader("Upload Allegro .xlsx", type=["xlsx","xls"], key="allegro")

# ── Helper functions ──
def fp(c): return PatternFill(fill_type='solid', start_color=c, end_color=c)
thin = Side(style='thin', color='CCCCCC')
bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
font_d = Font(name='Arial', size=9)
almt_l = Alignment(horizontal='left', vertical='center')

def cs(cell, bg, fc='000000', bold=False, size=9, halign='left', wrap=False, italic=False, num_fmt=None):
    cell.fill=fp(bg); cell.font=Font(name='Arial',bold=bold,italic=italic,color=fc,size=size)
    cell.alignment=Alignment(horizontal=halign,vertical='center',wrap_text=wrap); cell.border=bdr
    if num_fmt: cell.number_format=num_fmt

def mrow(ws,r,c1,c2,val,bg,fc='FFFFFF',bold=True,size=10,halign='center',wrap=False,italic=False,height=22):
    if c1!=c2: ws.merge_cells(start_row=r,start_column=c1,end_row=r,end_column=c2)
    cs(ws.cell(r,c1,val),bg,fc,bold,size,halign,wrap=wrap,italic=italic)
    ws.row_dimensions[r].height=height

def hlink(cell, sheet, bg, fc='004D40'):
    cell.hyperlink=Hyperlink(ref=f'{gcl(cell.column)}{cell.row}',location=f"'{sheet}'!A1")
    cell.font=Font(name='Arial',bold=True,color=fc,size=9,underline='single')
    cell.fill=fp(bg); cell.alignment=Alignment(horizontal='left',vertical='center'); cell.border=bdr

def clean_date(v):
    try:
        if v is None or (isinstance(v,float) and np.isnan(v)): return None
        return pd.to_datetime(v).date()
    except: return None

def categorise(row):
    tt=str(row.get('tradetype','')).upper()
    desc=str(row.get('description','')).lower()
    if tt=='TIME_CHARTER':   return 'Expected — Time Charter'
    if tt=='CASH_TRADE':     return 'Expected — Cash Trade'
    if 'prepayment' in desc: return 'Expected — Prepayment'
    if 'provisional' in desc:return 'Expected — Provisional Invoice'
    return 'Needs Investigation — Action Required'

FILLS = {'Not Invoiceable':'EDE7F6','Matched':'E8F5E9','Pending':'E3F2FD',
         'Unmatched Cash Flow':'FFF3E0','Unmatched Invoice Item':'FCE4EC','FFFFFF':'FFFFFF'}
SC_COLORS = {'Matched':'1B5E20','Pending':'1565C0','Not Invoiceable':'4A148C',
             'Unmatched Cash Flow':'E65100','Unmatched Invoice Item':'880E4F'}

def build_segment_sheet(sn, df, hc, title, note):
    cols=list(df.columns)
    wb=openpyxl.Workbook(); ws=wb.active; ws.title=sn[:31]; ws.sheet_view.showGridLines=False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
    cs(ws.cell(1,1,title),hc,'FFFFFF',True,12,'center'); ws.row_dimensions[1].height=30
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(cols))
    cs(ws.cell(2,1,note),'F5F5F5','424242',False,9,'left',wrap=True,italic=True); ws.row_dimensions[2].height=20
    for ci,c in enumerate(cols,1): cs(ws.cell(3,ci,c),hc,'FFFFFF',True,9,'center',wrap=True)
    ws.row_dimensions[3].height=32
    sc=cols.index('Invoice Item Match Status') if 'Invoice Item Match Status' in cols else None
    for ri,row in enumerate(df.values,4):
        fill_key=str(row[sc]) if sc is not None else 'FFFFFF'
        fill=fp(FILLS.get(fill_key,'FFFFFF')); ws.row_dimensions[ri].height=16
        for ci,val in enumerate(row,1):
            cell=ws.cell(ri,ci)
            cell.value=None if (val is None or (isinstance(val,float) and np.isnan(val))) else val
            cell.fill=fill; cell.font=font_d; cell.alignment=almt_l; cell.border=bdr
    for ci,h in enumerate(cols,1): ws.column_dimensions[gcl(ci)].width=min(max(len(str(h))+2,12),40)
    ws.freeze_panes='A4'
    buf=io.BytesIO(); wb.save(buf); buf.seek(0); return buf.read()

def build_ni_sheet(ni_rows, ni_cols):
    wb=Workbook(); ws=wb.active; ws.title='Needs Investigation'; ws.sheet_view.showGridLines=True
    col_labels={'jeragm_custom_inv_num':'Invoice Number','findetail':'Findetail','grossvalue':'Gross Value',
                'fee':'Fee','description':'Description','tradetype':'Trade Type','trade':'Trade',
                'position':'Position','creationdate':'Creation Date','Invoice Reason':'Reason'}
    col_w={'jeragm_custom_inv_num':26,'findetail':15,'grossvalue':14,'fee':14,'description':34,
           'tradetype':18,'trade':11,'position':11,'creationdate':13,'Invoice Reason':30}
    NUM_FORMATS={'findetail':'0','trade':'0','position':'0','fee':'0','grossvalue':'#,##0.00','creationdate':'YYYY-MM-DD'}
    for ci,col in enumerate(ni_cols,1): ws.column_dimensions[gcl(ci)].width=col_w.get(col,14)
    f_hdr=Font(name='Calibri',size=10,bold=True); f_data=Font(name='Calibri',size=10)
    a_left=Alignment(horizontal='left',vertical='center')
    ws.row_dimensions[1].height=18
    for ci,col in enumerate(ni_cols,1):
        c=ws.cell(1,ci,col_labels.get(col,col)); c.font=f_hdr; c.alignment=a_left
    prev_reason=None; ri=2
    for _,row_data in ni_rows.iterrows():
        reason=row_data['Invoice Reason']
        if reason!=prev_reason:
            inv_count=ni_rows[ni_rows['Invoice Reason']==reason]['jeragm_custom_inv_num'].nunique()
            ws.merge_cells(start_row=ri,start_column=1,end_row=ri,end_column=len(ni_cols))
            c=ws.cell(ri,1,f'{reason}  ({inv_count:,} invoices)')
            c.font=Font(name='Calibri',size=10,bold=True); c.alignment=a_left
            ws.row_dimensions[ri].height=15; ri+=1; prev_reason=reason
        ws.row_dimensions[ri].height=14
        for ci,col in enumerate(ni_cols,1):
            val=row_data[col]; cell=ws.cell(ri,ci)
            if val is None or (isinstance(val,float) and np.isnan(val)): cell.value=None
            elif col=='creationdate':
                try: cell.value=pd.to_datetime(val).date()
                except: cell.value=str(val)[:10] if val else None
            elif col in ('trade','position','findetail','fee'):
                try: cell.value=int(float(val))
                except: cell.value=val
            elif col=='grossvalue':
                try: cell.value=float(val)
                except: cell.value=val
            else: cell.value=val
            cell.font=f_data; cell.alignment=a_left
            if fmt:=NUM_FORMATS.get(col): cell.number_format=fmt
        ri+=1
    ws.freeze_panes='A2'
    buf=io.BytesIO(); wb.save(buf); buf.seek(0); return buf.read()

def run_recon(t_bytes, a_bytes, book, env, cutoff):
    progress = st.progress(0, text="Reading files...")

    t=pd.read_excel(io.BytesIO(t_bytes)); progress.progress(10, "Topaz loaded")
    a=pd.read_excel(io.BytesIO(a_bytes)); a.columns=[c.lower() for c in a.columns]
    a['creationdate']=pd.to_datetime(a['creationdate'],format='mixed')
    progress.progress(20, "Allegro loaded")

    # Match
    t['_mk']=pd.to_numeric(t['Invoice Item ID'],errors='coerce')
    a['_mk']=pd.to_numeric(a['findetail'],errors='coerce')
    a_r=a.rename(columns={c:f'Allegro_{c}' for c in a.columns})
    a_r['_mk']=pd.to_numeric(a_r['Allegro_findetail'],errors='coerce')
    merged=t.merge(a_r,on='_mk',how='left')
    merged.drop(columns=['_mk','Allegro_findetail'],inplace=True,errors='ignore')
    progress.progress(30, "Files matched")

    t_item=set(t['_mk'].dropna().astype(int))
    a_inv=set(a['jeragm_custom_inv_num'].dropna().unique())
    recv_invs=set(a[a['_mk'].isin(t_item)]['jeragm_custom_inv_num'])
    not_recv=a_inv-recv_invs
    a_dd=a.sort_values('creationdate').drop_duplicates('jeragm_custom_inv_num')
    nt=a_dd[a_dd['jeragm_custom_inv_num'].isin(not_recv)]
    arch=nt[nt['jeragm_archivedtrade']==1.0]; na_=nt[nt['jeragm_archivedtrade']!=1.0]
    exp=na_[na_['creationdate'].dt.date > cutoff.date()]
    ni_df=na_[na_['creationdate'].dt.date <= cutoff.date()]
    ni_invs=set(ni_df['jeragm_custom_inv_num'])
    ni_all_a=a[a['jeragm_custom_inv_num'].isin(ni_invs)]

    SENT=len(a_inv); RECV=len(recv_invs); NI_COUNT=len(ni_invs)
    progress.progress(40, "Computing NI categorisation...")

    # NI categorisation
    ni_rows=a[(a['jeragm_custom_inv_num'].isin(ni_invs)) & (~a['_mk'].isin(t_item))].copy()
    ni_rows['_reason']=ni_rows.apply(categorise,axis=1)
    priority={'Expected — Time Charter':1,'Expected — Cash Trade':2,'Expected — Prepayment':3,
              'Expected — Provisional Invoice':4,'Needs Investigation — Action Required':5}
    ni_rows['_pri']=ni_rows['_reason'].map(priority)
    inv_reason=(ni_rows.sort_values('_pri').drop_duplicates('jeragm_custom_inv_num')
                [['jeragm_custom_inv_num','_reason']].rename(columns={'_reason':'Invoice Reason'}))
    ni_rows=ni_rows.drop(columns=['_reason','_pri']).merge(inv_reason,on='jeragm_custom_inv_num')
    reason_order=['Needs Investigation — Action Required','Expected — Provisional Invoice',
                  'Expected — Time Charter','Expected — Prepayment','Expected — Cash Trade']
    ni_rows['_sort']=ni_rows['Invoice Reason'].map({r:i for i,r in enumerate(reason_order)})
    ni_rows=ni_rows.sort_values(['_sort','tradetype','trade','jeragm_custom_inv_num']).drop(columns=['_sort','_mk'])
    ni_cols=['jeragm_custom_inv_num','findetail','grossvalue','fee','description','tradetype',
             'trade','position','creationdate','Invoice Reason']
    ni_cols=[c for c in ni_cols if c in ni_rows.columns]
    ni_rows=ni_rows[ni_cols]
    summary_ni=(ni_rows.drop_duplicates('jeragm_custom_inv_num')
                .groupby('Invoice Reason').size().reindex(reason_order).fillna(0).astype(int))

    progress.progress(50, "Building segment sheets...")

    # Segments
    ts=t.groupby('Trade ID')['Invoice Item Match Status'].apply(lambda x: frozenset(x.dropna().unique())).reset_index()
    ts.columns=['Trade ID','Statuses']
    all_combos=ts['Statuses'].apply(lambda x: ' + '.join(sorted(x))).unique()

    seg_files={}
    for combo_str in sorted(all_combos):
        combo=frozenset(combo_str.split(' + '))
        ids=set(ts[ts['Statuses']==combo]['Trade ID'])
        df=merged[merged['Trade ID'].isin(ids)].sort_values(['Trade ID','Invoice Item Match Status'])
        sn=combo_str[:31]
        primary_status=sorted(combo)[0]
        hc={'Matched':'1B5E20','Pending':'1565C0','Not Invoiceable':'4A148C',
            'Unmatched Cash Flow':'E65100','Unmatched Invoice Item':'880E4F'}.get(primary_status,'37474F')
        seg_files[sn]=build_segment_sheet(sn,df,hc,f'{combo_str}  ({len(ids)} trades | {len(df)} rows)',combo_str)

    progress.progress(65, "Building All Data sheets...")

    CF_COL='Cash Flow Undiscounted (Native)'
    STATUS_COLS=['Matched','Pending','Not Invoiceable','Unmatched Cash Flow','Unmatched Invoice Item']

    # Matched All Data
    df_m=merged[merged['Invoice Item Match Status']=='Matched'].copy().sort_values('Trade ID')
    base=list(df_m.columns)
    extra=['Difference % of Invoice Item Amount','Invoice Amount Difference Category',
           'Topaz Invoice Date','Allegro Invoice Date','Invoice Date Match',
           'Topaz Invoice Settlement Day','Allegro Due Date','Due Date Match']
    all_cols=base+extra
    wb_m=io.BytesIO()
    wbx=xlsxwriter.Workbook(wb_m,{'constant_memory':True,'in_memory':True})
    ws_m=wbx.add_worksheet('Matched - All Data'); ws_m.hide_gridlines(2)
    hf=wbx.add_format({'bold':True,'font_name':'Arial','font_size':9,'font_color':'white','bg_color':'#1B5E20','align':'center','valign':'vcenter','text_wrap':True})
    df_fmt=wbx.add_format({'font_name':'Arial','font_size':9,'bg_color':'#E8F5E9','align':'left','valign':'vcenter'})
    pct_fmt=wbx.add_format({'font_name':'Arial','font_size':9,'bg_color':'#E8F5E9','num_format':'0.00%'})
    dhf=wbx.add_format({'bold':True,'font_name':'Arial','font_size':9,'font_color':'white','bg_color':'#0D47A1','align':'center','valign':'vcenter','text_wrap':True})
    mf=wbx.add_format({'font_name':'Arial','font_size':9,'bold':True,'font_color':'#1B5E20','bg_color':'#C8E6C9','align':'center'})
    xf=wbx.add_format({'font_name':'Arial','font_size':9,'bold':True,'font_color':'#B71C1C','bg_color':'#FFCDD2','align':'center'})
    ws_m.set_row(0,30)
    ws_m.merge_range(0,0,0,len(all_cols)-1,f'{book} {env} — Matched All Data  |  {len(df_m):,} rows',
        wbx.add_format({'bold':True,'font_name':'Arial','font_size':12,'font_color':'white','bg_color':'#1B5E20','align':'center','valign':'vcenter'}))
    ws_m.set_row(1,20)
    ws_m.merge_range(1,0,1,len(all_cols)-1,'ALL Matched rows. ABS(V-Z)<0.01 = Exact Match.',
        wbx.add_format({'italic':True,'font_name':'Arial','font_size':9,'font_color':'#424242','bg_color':'#F5F5F5','align':'left','valign':'vcenter'}))
    ws_m.set_row(2,36)
    date_cols={'Topaz Invoice Date','Allegro Invoice Date','Invoice Date Match',
               'Topaz Invoice Settlement Day','Allegro Due Date','Due Date Match'}
    for ci,col in enumerate(all_cols): ws_m.write(2,ci,col,dhf if col in date_cols else hf)
    cf_i=base.index(CF_COL) if CF_COL in base else None
    ia_i=base.index('Invoice Item Amount') if 'Invoice Item Amount' in base else None
    td_i=base.index('Invoice Date') if 'Invoice Date' in base else None
    isd_i=base.index('Invoice Settlement Day') if 'Invoice Settlement Day' in base else None
    aid_i=base.index('Allegro_invoicedate') if 'Allegro_invoicedate' in base else None
    asd_i=base.index('Allegro_duedate') if 'Allegro_duedate' in base else None
    for ri,row in enumerate(df_m.values,3):
        ws_m.set_row(ri,15)
        for ci,val in enumerate(row): ws_m.write(ri,ci,None if (isinstance(val,float) and np.isnan(val)) else val,df_fmt)
        if cf_i is not None and ia_i is not None:
            try:
                v_n=float(row[cf_i]); z_n=float(row[ia_i])
                if abs(v_n-z_n)<0.01: cat='Exact Match'; pct=0.0
                elif v_n==-z_n and v_n!=0: cat='Sign Difference'; pct=abs(v_n-z_n)/abs(z_n) if z_n!=0 else 0
                else: cat='Amount Difference'; pct=abs(v_n-z_n)/abs(z_n) if z_n!=0 else 0
            except: cat=''; pct=0
        else: cat=''; pct=0
        ws_m.write(ri,len(base),pct,pct_fmt); ws_m.write(ri,len(base)+1,cat,df_fmt)
        def gd(idx):
            if idx is None: return None
            try: return pd.to_datetime(row[idx]).date() if row[idx] else None
            except: return None
        t_inv=gd(td_i); a_inv=gd(aid_i); t_set=gd(isd_i); a_set=gd(asd_i)
        ws_m.write(ri,len(base)+2,str(t_inv) if t_inv else None,df_fmt)
        ws_m.write(ri,len(base)+3,str(a_inv) if a_inv else None,df_fmt)
        if t_inv and a_inv:
            ok=t_inv==a_inv; ws_m.write(ri,len(base)+4,'✓ Match' if ok else '✗ Mismatch',mf if ok else xf)
        else: ws_m.write(ri,len(base)+4,'—',df_fmt)
        ws_m.write(ri,len(base)+5,str(t_set) if t_set else None,df_fmt)
        ws_m.write(ri,len(base)+6,str(a_set) if a_set else None,df_fmt)
        if t_set and a_set:
            ok=t_set==a_set; ws_m.write(ri,len(base)+7,'✓ Match' if ok else '✗ Mismatch',mf if ok else xf)
        else: ws_m.write(ri,len(base)+7,'—',df_fmt)
    for ci,h in enumerate(all_cols): ws_m.set_column(ci,ci,min(max(len(str(h))+2,12),40))
    ws_m.freeze_panes(3,0); wbx.close(); wb_m.seek(0)
    matched_all_bytes=wb_m.read()

    # Other All Data
    all_data_files={}
    for status,fill_color,hc in [
        ('Pending','E3F2FD','1565C0'),('Not Invoiceable','EDE7F6','4A148C'),
        ('Unmatched Cash Flow','FFF3E0','E65100'),('Unmatched Invoice Item','FCE4EC','880E4F')]:
        df2=merged[merged['Invoice Item Match Status']==status].copy().sort_values('Trade ID')
        if len(df2)==0: continue
        c_list=list(df2.columns); fill=fp(fill_color)
        wb2=openpyxl.Workbook(); ws2=wb2.active; ws2.title=f'{status[:20]}-All'; ws2.sheet_view.showGridLines=False
        ws2.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(c_list))
        c=ws2.cell(1,1,f'{book} {env} — {status} All  |  {len(df2):,} rows')
        c.fill=fp(hc); c.font=Font(name='Arial',bold=True,color='FFFFFF',size=12)
        c.alignment=Alignment(horizontal='center',vertical='center'); c.border=bdr; ws2.row_dimensions[1].height=30
        ws2.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(c_list))
        c2=ws2.cell(2,1,f'All {status} rows.'); c2.fill=fp('F5F5F5')
        c2.font=Font(name='Arial',italic=True,color='424242',size=9); c2.alignment=Alignment(horizontal='left',vertical='center'); c2.border=bdr; ws2.row_dimensions[2].height=20
        for ci,col in enumerate(c_list,1):
            c=ws2.cell(3,ci,col); c.fill=fp(hc); c.font=Font(name='Arial',bold=True,color='FFFFFF',size=9)
            c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=bdr
        ws2.row_dimensions[3].height=32
        for ri,row in enumerate(df2.values,4):
            ws2.row_dimensions[ri].height=16
            for ci,val in enumerate(row,1):
                cell=ws2.cell(ri,ci)
                cell.value=None if (val is None or (isinstance(val,float) and np.isnan(val))) else val
                cell.fill=fill; cell.font=font_d; cell.alignment=almt_l; cell.border=bdr
        for ci,h in enumerate(c_list,1): ws2.column_dimensions[gcl(ci)].width=min(max(len(str(h))+2,12),40)
        ws2.freeze_panes='A4'
        buf=io.BytesIO(); wb2.save(buf); buf.seek(0)
        all_data_files[f'{status} - All Data']=buf.read()

    progress.progress(80, "Building Summary sheet...")

    # Summary
    topaz_nums=set()
    for tid in t['Trade ID'].dropna().unique():
        parts=tid.split('-')
        if len(parts)>=2:
            try: topaz_nums.add(int(parts[1]))
            except: pass
    ni_trades_all=ni_all_a['trade'].dropna().astype(int).unique()
    not_found=[]; ni_only=[]; mixed=[]; combo_r={}
    for tr in ni_trades_all:
        if tr not in topaz_nums: not_found.append(tr)
        else:
            m=t[t['Trade ID'].str.contains(str(tr),na=False)]
            st=frozenset(m['Invoice Item Match Status'].dropna().unique())
            if st==frozenset({'Not Invoiceable'}): ni_only.append(tr)
            else:
                mixed.append(tr); combo=' + '.join(sorted(st))
                invs=set(ni_all_a[ni_all_a['trade']==tr]['jeragm_custom_inv_num'])
                combo_r.setdefault(combo,{'trades':0,'inv_set':set()})
                combo_r[combo]['trades']+=1; combo_r[combo]['inv_set'].update(invs)
    def ic(tl): return len(set(ni_all_a[ni_all_a['trade'].isin(tl)]['jeragm_custom_inv_num']))

    pivot=t.groupby(['Trade Type','Invoice Item Match Status']).size().unstack(fill_value=0)
    for col in STATUS_COLS:
        if col not in pivot.columns: pivot[col]=0
    pivot=pivot[STATUS_COLS]; pivot.loc['TOTAL']=pivot.sum()
    TOTAL_ROWS=len(t)
    not_recv_count=len(not_recv)
    NI_TRADES=len(ni_trades_all)

    wb_s=openpyxl.Workbook(); ws_s=wb_s.active; ws_s.title='Summary'; ws_s.sheet_view.showGridLines=False
    for ci,w in {1:28,2:10,3:11,4:10,5:36,6:10,7:10,8:10,9:36}.items(): ws_s.column_dimensions[gcl(ci)].width=w
    R=1
    mrow(ws_s,R,1,9,f'{book} BOOK — Invoice Reconciliation Summary ({env})','004D40','FFFFFF',True,14,'center',height=36); R+=1
    mrow(ws_s,R,1,9,f'Topaz vs Allegro  |  Book: {book}  |  {env}  |  Cutoff: {cutoff_str}  |  Match: Invoice Item ID ↔ findetail',
         '00695C','FFFFFF',False,9,'center',italic=True,height=18); R+=2
    mrow(ws_s,R,1,9,'INVOICE FLOW — Allegro Sent  →  Topaz Received','004D40','FFFFFF',True,11,'center',height=26); R+=1
    ws_s.row_dimensions[R].height=18
    for c1,c2,h,bg in [(1,1,'','37474F'),(2,2,'Count','37474F'),(3,3,'% of Sent','37474F'),(4,9,'Description','37474F')]:
        ws_s.merge_cells(start_row=R,start_column=c1,end_row=R,end_column=c2); cs(ws_s.cell(R,c1,h),bg,'FFFFFF',True,9,'center')
    R+=1
    for lbl,count,pct_val,hc2,rc2,desc,green in [
        ('Allegro Sent',SENT,1.0,'006064','E0F2F1','Total distinct Custom Invoice Numbers',False),
        ('Topaz Received',RECV,RECV/SENT if SENT else 0,'1B5E20','E8F5E9','Matched via Invoice Item ID ↔ findetail',True),
    ]:
        ws_s.row_dimensions[R].height=24; cs(ws_s.cell(R,1,lbl),hc2,'FFFFFF',True,11,'center')
        c=ws_s.cell(R,2,count); cs(c,rc2,hc2,True,13,'center',num_fmt='#,##0')
        c2x=ws_s.cell(R,3,pct_val); cs(c2x,'00FF00' if green else rc2,hc2,True,11,'center',num_fmt='0.0%')
        ws_s.merge_cells(start_row=R,start_column=4,end_row=R,end_column=9)
        cs(ws_s.cell(R,4,desc),rc2,'004D40',False,9,'left',wrap=True); R+=1
    R+=1
    mrow(ws_s,R,1,9,f'NOT IN TOPAZ — {not_recv_count:,} Custom Invoice Numbers','B71C1C','FFFFFF',True,11,'center',height=26); R+=1
    ws_s.row_dimensions[R].height=18
    for c1,c2,h,bg in [(1,1,'Reason','37474F'),(2,2,'Count','37474F'),(3,3,'% of Not Recv','37474F'),(4,9,'Explanation','37474F')]:
        ws_s.merge_cells(start_row=R,start_column=c1,end_row=R,end_column=c2); cs(ws_s.cell(R,c1,h),bg,'FFFFFF',True,9,'center')
    R+=1
    for reason,count,hc2,rc2,desc in [
        ('Archived',len(arch),'546E7A','ECEFF1','jeragm_archivedtrade = 1'),
        ('Expected — Created after cutoff',len(exp),'1565C0','E3F2FD','Created after cutoff date'),
        ('Needs Investigation',NI_COUNT,'E65100','FFF3E0','⚠  NOT archived + on or before cutoff'),
    ]:
        if count==0: continue
        ws_s.row_dimensions[R].height=24; cs(ws_s.cell(R,1,reason),hc2,'FFFFFF',True,10,'center',wrap=True)
        c=ws_s.cell(R,2,count); cs(c,rc2,hc2,True,13,'center',num_fmt='#,##0')
        c3=ws_s.cell(R,3,count/not_recv_count if not_recv_count else 0); cs(c3,rc2,hc2,False,10,'center',num_fmt='0.0%')
        ws_s.merge_cells(start_row=R,start_column=4,end_row=R,end_column=9)
        cs(ws_s.cell(R,4,desc),rc2,'B71C1C' if '⚠' in desc else '37474F',False,9,'left',wrap=True); R+=1
    ws_s.row_dimensions[R].height=22; cs(ws_s.cell(R,1,'Total Not in Topaz'),'263238','FFFFFF',True,10,'center')
    c=ws_s.cell(R,2,not_recv_count); cs(c,'263238','FFFFFF',True,13,'center',num_fmt='#,##0')
    c3=ws_s.cell(R,3,1.0); cs(c3,'263238','FFFFFF',True,10,'center',num_fmt='0%')
    ws_s.merge_cells(start_row=R,start_column=4,end_row=R,end_column=9); cs(ws_s.cell(R,4,''),'263238'); R+=2
    mrow(ws_s,R,1,9,f'NEEDS INVESTIGATION — {NI_COUNT:,} Custom Invoice Numbers  (Trade-level Analysis)','E65100','FFFFFF',True,11,'center',height=26); R+=1
    ws_s.merge_cells(start_row=R,start_column=1,end_row=R,end_column=9)
    n=ws_s.cell(R,1,f'ALL {NI_COUNT:,} invoice numbers NOT in Topaz. {NI_TRADES:,} distinct Allegro trades checked:')
    n.fill=fp('FFF3E0'); n.font=Font(name='Arial',size=9,italic=True,color='E65100')
    n.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True); n.border=bdr; ws_s.row_dimensions[R].height=18; R+=1
    ws_s.row_dimensions[R].height=22
    for c1,c2,h,bg in [(1,1,'Category','37474F'),(2,2,'Trades','37474F'),(3,3,'Invoice Numbers','37474F'),(4,6,'Description','37474F'),(7,9,'Sheet Reference','004D40')]:
        ws_s.merge_cells(start_row=R,start_column=c1,end_row=R,end_column=c2); cs(ws_s.cell(R,c1,h),bg,'FFFFFF',True,9,'center')
    R+=1
    for lbl,trades,invs,hc2,rc2,desc in [
        ('Trade NOT Found in Topaz',len(not_found),ic(not_found),'B71C1C','FFCDD2',f'{len(not_found)} trades completely absent'),
        ('Trade in Topaz — Not Invoiceable Only',len(ni_only),ic(ni_only),'4A148C','EDE7F6',f'{len(ni_only)} trades all Not Invoiceable'),
        ('Trade in Topaz — Mixed Statuses',len(mixed),ic(mixed),'00695C','E0F2F1',f'{len(mixed)} trades with mixed statuses'),
    ]:
        ws_s.row_dimensions[R].height=24; cs(ws_s.cell(R,1,lbl),hc2,'FFFFFF',True,9,'center',wrap=True)
        c=ws_s.cell(R,2,trades); cs(c,rc2,hc2,True,12,'center',num_fmt='#,##0')
        c=ws_s.cell(R,3,invs);   cs(c,rc2,hc2,False,10,'center',num_fmt='#,##0')
        ws_s.merge_cells(start_row=R,start_column=4,end_row=R,end_column=6)
        cs(ws_s.cell(R,4,desc),rc2,'37474F',False,9,'left',wrap=True)
        ws_s.merge_cells(start_row=R,start_column=7,end_row=R,end_column=9)
        cs(ws_s.cell(R,7,''),'E0F2F1'); R+=1
    ws_s.row_dimensions[R].height=22; cs(ws_s.cell(R,1,f'TOTAL — {NI_TRADES:,} trades'),'263238','FFFFFF',True,10,'center')
    c=ws_s.cell(R,2,NI_TRADES); cs(c,'263238','FFFFFF',True,13,'center',num_fmt='#,##0')
    c=ws_s.cell(R,3,NI_COUNT);  cs(c,'263238','FFFFFF',True,10,'center',num_fmt='#,##0')
    ws_s.merge_cells(start_row=R,start_column=4,end_row=R,end_column=9); cs(ws_s.cell(R,4,''),'263238'); R+=2
    mrow(ws_s,R,1,9,'TRADE TYPE BREAKDOWN','37474F','FFFFFF',True,11,'center',height=26); R+=1
    ws_s.row_dimensions[R].height=34; cs(ws_s.cell(R,1,'Trade Type'),'37474F','FFFFFF',True,9,'left')
    for ci,h in enumerate(STATUS_COLS,2): cs(ws_s.cell(R,ci,h),SC_COLORS[h],'FFFFFF',True,9,'center',wrap=True)
    cs(ws_s.cell(R,7,'TOTAL'),'263238','FFFFFF',True,9,'center')
    ws_s.merge_cells(start_row=R,start_column=8,end_row=R,end_column=9); cs(ws_s.cell(R,8,''),'263238'); R+=1
    for i,(tt,rd) in enumerate(pivot.iterrows()):
        is_t=(tt=='TOTAL'); bg2='263238' if is_t else ('F0F4F8' if i%2==0 else 'FFFFFF'); fc2='FFFFFF' if is_t else '000000'
        ws_s.row_dimensions[R].height=18; cs(ws_s.cell(R,1,tt),bg2,fc2,is_t,9,'left'); rt=0
        for ci,s in enumerate(STATUS_COLS,2):
            val=int(rd.get(s,0)); rt+=val; cell=ws_s.cell(R,ci)
            if val>0: cs(cell,bg2,SC_COLORS[s] if not is_t else 'FFFFFF',not is_t,9,'center',num_fmt='#,##0'); cell.value=val
            else: cs(cell,bg2,'999999' if not is_t else fc2,False,9,'center'); cell.value='—'
        cell=ws_s.cell(R,7,rt); cs(cell,bg2,'FFFFFF' if is_t else '37474F',True,10,'center',num_fmt='#,##0')
        ws_s.merge_cells(start_row=R,start_column=8,end_row=R,end_column=9); cs(ws_s.cell(R,8,''),bg2); R+=1
    buf_s=io.BytesIO(); wb_s.save(buf_s); buf_s.seek(0)
    summary_bytes=buf_s.read()

    progress.progress(90, "Assembling final workbook...")

    # NI sheet
    ni_bytes=build_ni_sheet(ni_rows,ni_cols)

    # Assemble
    with zipfile.ZipFile(io.BytesIO(summary_bytes),'r') as z:
        files={n:z.read(n) for n in z.namelist()}
    styles=files['xl/styles.xml']
    rel_xml=files['xl/_rels/workbook.xml.rels'].decode('utf-8')
    rel_xml=rel_xml.replace('Target="/xl/worksheets/','Target="worksheets/')
    existing=re.findall(r'sheet(\d+)\.xml',' '.join(files.keys()))
    next_sn=max(int(x) for x in existing)+1
    rids=re.findall(r'rId(\d+)',rel_xml)
    next_rid=max(int(x) for x in rids)+1
    sheet_tags='<sheet name="Summary" sheetId="1" r:id="rId1"/>'

    all_sheets=[
        *[(k,v) for k,v in seg_files.items()],
        ('Matched - All Data', matched_all_bytes),
        *[(k,v) for k,v in all_data_files.items()],
        ('Needs Investigation', ni_bytes),
    ]

    for sn,data in all_sheets:
        with zipfile.ZipFile(io.BytesIO(data),'r') as z:
            files[f'xl/worksheets/sheet{next_sn}.xml']=z.read('xl/worksheets/sheet1.xml')
        sn_esc=sn.replace('&','&amp;')
        sheet_tags+=f'<sheet name="{sn_esc}" sheetId="{next_sn}" r:id="rId{next_rid}"/>'
        rel_xml=rel_xml.replace('</Relationships>',
            f'<Relationship Id="rId{next_rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{next_sn}.xml"/></Relationships>')
        next_sn+=1; next_rid+=1

    wb_xml=('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<workbookPr/><bookViews><workbookView/></bookViews>'
        f'<sheets>{sheet_tags}</sheets>'
        '<calcPr calcId="124519" fullCalcOnLoad="1"/>'
        '</workbook>')
    files['xl/styles.xml']=styles
    files['xl/workbook.xml']=wb_xml.encode('utf-8')
    files['xl/_rels/workbook.xml.rels']=rel_xml.encode('utf-8')

    # Fix alpha + strip NI styles
    st2=files['xl/styles.xml'].decode('utf-8')
    files['xl/styles.xml']=re.sub(r'rgb="00([0-9A-Fa-f]{6})"',r'rgb="FF\1"',st2).encode('utf-8')
    sheets_map=dict(re.findall(r'<sheet name="([^"]+)" sheetId="(\d+)"',files['xl/workbook.xml'].decode()))
    ni_sn=sheets_map.get('Needs Investigation')
    if ni_sn:
        ni_xml=files[f'xl/worksheets/sheet{ni_sn}.xml'].decode('utf-8')
        files[f'xl/worksheets/sheet{ni_sn}.xml']=re.sub(r' s="\d+"','',ni_xml).encode('utf-8')

    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for fn,d in files.items(): z.writestr(fn,d)
    out.seek(0)
    progress.progress(100, "Done!")
    return out.read(), SENT, RECV, NI_COUNT, summary_ni

# ── Run ──
if topaz_file and allegro_file:
    st.markdown("---")
    if st.button("▶  Run Reconciliation", type="primary", use_container_width=True):
        with st.spinner("Running reconciliation..."):
            try:
                t_bytes=topaz_file.read(); a_bytes=allegro_file.read()
                result_bytes,SENT,RECV,NI,summary_ni=run_recon(t_bytes,a_bytes,book,env,cutoff)

                st.success("✅ Reconciliation complete!")

                c1,c2,c3,c4=st.columns(4)
                c1.metric("Allegro Sent",f"{SENT:,}")
                c2.metric("Topaz Received",f"{RECV:,}",f"{RECV/SENT*100:.1f}%" if SENT else "")
                c3.metric("Not Matched",f"{SENT-RECV:,}")
                c4.metric("Needs Investigation",f"{NI:,}")

                st.markdown("**NI Breakdown:**")
                ni_df=pd.DataFrame(summary_ni).reset_index()
                ni_df.columns=['Reason','Count']
                ni_df=ni_df[ni_df['Count']>0]
                st.dataframe(ni_df,use_container_width=True,hide_index=True)

                fname=f"{book}_{env}_Invoice_Reconciliation.xlsx"
                st.download_button(
                    label=f"⬇  Download {fname}",
                    data=result_bytes,
                    file_name=fname,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    type="primary"
                )
            except Exception as e:
                st.error(f"Error: {e}")
                import traceback; st.code(traceback.format_exc())
else:
    st.info("Upload both Topaz and Allegro files to begin.")
